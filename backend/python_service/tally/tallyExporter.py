import sys
import json
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from datetime import datetime
import xml.etree.ElementTree as ET

import importlib.util

# Import the existing parser dynamically to avoid conflict with standard library 'parser'
parser_path = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "parser.py")
spec = importlib.util.spec_from_file_location("local_parser", parser_path)
local_parser = importlib.util.module_from_spec(spec)
spec.loader.exec_module(local_parser)

def generate_excel(transactions, output_path, metadata):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Transaction"

    # Define styles
    font_bold = Font(bold=True)
    font_small = Font(size=9)
    fill_yellow = PatternFill(start_color="FFFFFF00", end_color="FFFFFF00", fill_type="solid")
    align_center = Alignment(horizontal="center", vertical="center")
    
    # Write top headers (like HDFC Bank template)
    ws['A1'] = metadata.get('bankName', 'BANK Ltd.')
    ws['A1'].font = font_bold
    ws['D1'] = "Page No. : 1"
    ws['F1'] = "Statement of accounts"

    ws['A4'] = "M/S. KNOWLEDGE MOVERS AND SHAKERS" # Example placeholder, should be from metadata if available
    
    ws['H6'] = f"Account Branch : {metadata.get('branch', '')}"
    ws['H12'] = f"Account No : {metadata.get('accountNumber', '')}"
    
    # Headers
    headers = [
        "Date", "Narration", "Chq./Ref.No.", "Value Dt", 
        "Withdrawal Amt.", "Deposit Amt.", "Closing Balance", 
        "Remarks - 1", "Remarks - 2", "Remarks - 3", "Remarks - 4"
    ]
    
    row_num = 18
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=row_num, column=col_num)
        cell.value = header
        cell.font = font_bold
        cell.alignment = align_center

    # Sub-headers (Debit/Credit)
    ws.cell(row=19, column=5).value = "Debit"
    ws.cell(row=19, column=6).value = "Credit"

    row_num = 20
    total_debit = 0
    total_credit = 0
    
    for txn in transactions:
        amount_val = float(txn.get('amount', 0))
        debit = amount_val if txn.get('type') == 'debit' else ""
        credit = amount_val if txn.get('type') == 'credit' else ""
        
        if debit: total_debit += debit
        if credit: total_credit += credit

        # Highlight random rows for testing (or base it on a condition)
        # We'll just leave it plain for now, as user will fill it
        
        ws.cell(row=row_num, column=1).value = txn.get('date', '')
        ws.cell(row=row_num, column=2).value = txn.get('description', '')
        ws.cell(row=row_num, column=3).value = txn.get('refNo', '')
        ws.cell(row=row_num, column=4).value = txn.get('date', '')
        ws.cell(row=row_num, column=5).value = debit
        ws.cell(row=row_num, column=6).value = credit
        ws.cell(row=row_num, column=7).value = txn.get('closingBalance', '')
        
        row_num += 1

    # Statement Summary
    row_num += 2
    ws.cell(row=row_num, column=1).value = "STATEMENT SUMMARY :-"
    ws.cell(row=row_num, column=1).font = font_bold
    
    row_num += 1
    ws.cell(row=row_num, column=1).value = "Opening Balance"
    ws.cell(row=row_num, column=5).value = "Debits"
    ws.cell(row=row_num, column=6).value = "Credits"
    ws.cell(row=row_num, column=7).value = "Closing Bal"
    
    row_num += 1
    ws.cell(row=row_num, column=5).value = total_debit
    ws.cell(row=row_num, column=6).value = total_credit
    
    # Set column widths
    ws.column_dimensions['A'].width = 12
    ws.column_dimensions['B'].width = 50
    ws.column_dimensions['C'].width = 20
    ws.column_dimensions['D'].width = 12
    ws.column_dimensions['E'].width = 15
    ws.column_dimensions['F'].width = 15
    ws.column_dimensions['G'].width = 15
    ws.column_dimensions['H'].width = 15

    wb.save(output_path)
    return output_path

def generate_xml(transactions, output_path):
    envelope = ET.Element('ENVELOPE')
    header = ET.SubElement(envelope, 'HEADER')
    tallyreq = ET.SubElement(header, 'TALLYREQUEST')
    tallyreq.text = "Import Data"
    
    body = ET.SubElement(envelope, 'BODY')
    import_data = ET.SubElement(body, 'IMPORTDATA')
    req_desc = ET.SubElement(import_data, 'REQUESTDESC')
    report_name = ET.SubElement(req_desc, 'REPORTNAME')
    report_name.text = 'Vouchers'
    static_vars = ET.SubElement(req_desc, 'STATICVARIABLES')
    svc_cmp = ET.SubElement(static_vars, 'SVCURRENTCOMPANY')
    svc_cmp.text = 'My Company'
    
    req_data = ET.SubElement(import_data, 'REQUESTDATA')
    
    for txn in transactions:
        tally_msg = ET.SubElement(req_data, 'TALLYMESSAGE', {'xmlns:UDF': 'TallyUDF'})
        voucher = ET.SubElement(tally_msg, 'VOUCHER', {'VCHTYPE': 'Receipt' if txn.get('type') == 'credit' else 'Payment', 'ACTION': 'Create'})
        
        date_el = ET.SubElement(voucher, 'DATE')
        try:
            date_obj = datetime.strptime(txn.get('date'), '%d/%m/%Y')
            date_el.text = date_obj.strftime('%Y%m%d')
        except:
            date_el.text = "20260401"
        
        narration = ET.SubElement(voucher, 'NARRATION')
        narration.text = txn.get('description', '')
        
        amount_val = float(txn.get('amount', 0))
        
        # Ledger entries
        led_entry1 = ET.SubElement(voucher, 'ALLLEDGERENTRIES.LIST')
        led_name1 = ET.SubElement(led_entry1, 'LEDGERNAME')
        led_name1.text = 'Bank Account'
        is_deemed1 = ET.SubElement(led_entry1, 'ISDEEMEDPOSITIVE')
        is_deemed1.text = 'Yes' if txn.get('type') == 'credit' else 'No'
        amt1 = ET.SubElement(led_entry1, 'AMOUNT')
        amt1.text = str(-amount_val) if txn.get('type') == 'credit' else str(amount_val)
        
        led_entry2 = ET.SubElement(voucher, 'ALLLEDGERENTRIES.LIST')
        led_name2 = ET.SubElement(led_entry2, 'LEDGERNAME')
        led_name2.text = 'Suspense A/c'
        is_deemed2 = ET.SubElement(led_entry2, 'ISDEEMEDPOSITIVE')
        is_deemed2.text = 'No' if txn.get('type') == 'credit' else 'Yes'
        amt2 = ET.SubElement(led_entry2, 'AMOUNT')
        amt2.text = str(amount_val) if txn.get('type') == 'credit' else str(-amount_val)
        
    tree = ET.ElementTree(envelope)
    tree.write(output_path, encoding='utf-8', xml_declaration=True)
    return output_path

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(json.dumps({"error": "No PDF provided."}))
        sys.exit(1)
        
    pdf_path = sys.argv[1]
    
    from io import StringIO
    old_stdout = sys.stdout
    sys.stdout = mystdout = StringIO()
    
    try:
        local_parser.parse_pdf(pdf_path)
    except Exception as e:
        sys.stdout = old_stdout
        print(json.dumps({"error": str(e)}))
        sys.exit(1)
        
    sys.stdout = old_stdout
    parser_output = mystdout.getvalue()
    
    try:
        parsed_data = json.loads(parser_output)
    except Exception as e:
        if "error" in parser_output:
            print(parser_output)
        else:
            print(json.dumps({"error": "Failed to parse transactions", "details": parser_output}))
        sys.exit(1)

    if isinstance(parsed_data, dict) and "error" in parsed_data:
        print(json.dumps(parsed_data))
        sys.exit(1)
        
    # The parser returns either a list of transactions, or { transactions: [...], metadata: {...} }
    transactions = parsed_data.get('transactions', []) if isinstance(parsed_data, dict) else parsed_data
    metadata = parsed_data.get('metadata', {}) if isinstance(parsed_data, dict) else {}
        
    base_name = os.path.basename(pdf_path)
    base_name_no_ext = os.path.splitext(base_name)[0]
    upload_dir = os.path.dirname(pdf_path)
    
    excel_filename = f"{base_name_no_ext}_Tally.xlsx"
    xml_filename = f"{base_name_no_ext}_Tally.xml"
    
    excel_path = os.path.join(upload_dir, excel_filename)
    xml_path = os.path.join(upload_dir, xml_filename)
    
    try:
        generate_excel(transactions, excel_path, metadata)
        generate_xml(transactions, xml_path)
        
        result = {
            "success": True,
            "excelUrl": f"/uploads/{excel_filename}",
            "xmlUrl": f"/uploads/{xml_filename}",
            "transactionCount": len(transactions)
        }
        print(json.dumps(result))
    except Exception as e:
        print(json.dumps({"error": f"Failed to generate files: {str(e)}"}))
        sys.exit(1)
